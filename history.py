"""
历史汇总模块

职责：
  - 扫描所有历史 tradeable_opportunities.csv，提取每日 top5×2 记录
  - 生成 analysis/top_opportunities_history.xlsx（含日期下拉 + 历史数据两个 sheet）
"""

import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from paths import ANALYSIS_ROOT, HISTORY_XLSX

# ── 历史汇总相关常量 ──────────────────────────────────────────────────────────
# 历史数据 sheet 列头（按顺序）
_HIST_HEADERS = [
    "日期", "冻结天数", "排名",
    "商品名", "分类", "买入天数", "卖出天数", "天数差",
    "买入价(最低)", "卖出价(最低)", "税后利润", "收益率%",
    "买入价(均价)", "卖出价(均价)", "税后利润(均)", "收益率%(均)",
]
# tradeable_opportunities.csv 中对应列（前 3 个由代码填充）
_HIST_SRC = [
    None, None, None,
    "商品名", "分类", "买入天数", "卖出天数", "天数差",
    "买入价(最低)", "卖出价(最低)", "利润_min", "收益率%_min",
    "买入价(均价)", "卖出价(均价)", "利润_avg", "收益率%_avg",
]
# Sheet 1 展示列（去掉日期/冻结天数/排名），及对应 历史数据 sheet 列字母
_DISP_HEADERS = _HIST_HEADERS[3:]   # 13 列
_HIST_DATA_COLS = [get_column_letter(i) for i in range(4, 4 + len(_DISP_HEADERS))]  # D~P
# ─────────────────────────────────────────────────────────────────────────────


def _collect_all_tops() -> list:
    """
    扫描 analysis/XXXXXX/ 下所有 tradeable_opportunities.csv，
    提取每日 top5×2（30天/7天冻结），加上日期、冻结天数、排名列。
    返回按 日期→冻结天数→排名 排序的字典列表。
    """
    all_rows = []
    if not os.path.isdir(ANALYSIS_ROOT):
        return all_rows
    for entry in sorted(os.listdir(ANALYSIS_ROOT)):
        if not re.fullmatch(r"\d{8}", entry):
            continue
        csv_path = os.path.join(ANALYSIS_ROOT, entry, "tradeable_opportunities.csv")
        if not os.path.isfile(csv_path):
            continue
        try:
            df = pd.read_csv(csv_path, encoding="utf-8-sig")
        except Exception as e:
            print(f"  [警告] 跳过 {entry}: {e}")
            continue

        date_int = int(entry)
        for freeze in [30, 7]:
            subset = (df[df["冻结天数"] == freeze]
                      .sort_values("收益率%_min", ascending=False)
                      .head(5))
            for rank, (_, row) in enumerate(subset.iterrows(), start=1):
                record = {"日期": date_int, "冻结天数": freeze, "排名": rank}
                for src_col, hist_hdr in zip(_HIST_SRC[3:], _HIST_HEADERS[3:]):
                    val = row.get(src_col, "")
                    record[hist_hdr] = "" if pd.isna(val) else val
                all_rows.append(record)
    return all_rows


def build_history_xlsx():
    """
    生成 analysis/top_opportunities_history.xlsx，含两个 sheet：
      - 套利速览：日期下拉选择 + SUMPRODUCT 公式动态展示当日 top5×2
      - 历史数据：所有历史记录原始行
    """
    rows = _collect_all_tops()
    if not rows:
        print("  [警告] 未找到任何历史数据，跳过历史汇总。")
        return

    all_dates = sorted({r["日期"] for r in rows})
    latest_date = all_dates[-1]

    # ── 配色（与 export_xlsx 保持一致） ───────────────────────────────────────
    CLR_30_TITLE   = "B8860B"; CLR_30_HEADER  = "FFD700"
    CLR_30_ROW_A   = "FFFACD"; CLR_30_ROW_B   = "FFF8DC"
    CLR_7_TITLE    = "1B4F8A"; CLR_7_HEADER   = "4472C4"
    CLR_7_ROW_A    = "DCE6F1"; CLR_7_ROW_B    = "EEF4FB"
    CLR_PROFIT_POS = "1E8449"; CLR_PROFIT_NEG = "C0392B"
    CLR_FONT_DARK  = "1A1A1A"; CLR_FONT_WHITE = "FFFFFF"

    thin   = Side(style="thin",   color="BFBFBF")
    thick  = Side(style="medium", color="888888")
    b_all  = Border(left=thin,  right=thin,  top=thin, bottom=thin)
    b_sep  = Border(left=thick, right=thin,  top=thin, bottom=thin)

    def fill(hex_): return PatternFill("solid", fgColor=hex_)
    def myfont(bold=False, size=10, color=CLR_FONT_DARK):
        return Font(name="微软雅黑", bold=bold, color=color, size=size)

    center = Alignment(horizontal="center", vertical="center")
    left_  = Alignment(horizontal="left",   vertical="center")

    wb = Workbook()

    # ════════════════════════════════════════════════════════════════════════
    # Sheet 1: 套利速览
    # ════════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "套利速览"

    n_disp = len(_DISP_HEADERS)   # 13
    COL_WIDTHS = [12, 10, 8, 8, 7, 13, 13, 11, 9, 13, 13, 11, 9]
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ── 日期选择行 ────────────────────────────────────────────────────────
    label = ws1["A1"]
    label.value = "选择日期："
    label.fill  = fill("E8F0FE")
    label.font  = myfont(bold=True, size=11)
    label.alignment = Alignment(horizontal="right", vertical="center")
    label.border = b_all

    ws1.merge_cells("B1:C1")
    date_cell = ws1["B1"]
    date_cell.value     = latest_date
    date_cell.fill      = fill("FFFFFF")
    date_cell.font      = myfont(bold=True, size=12, color="1B4F8A")
    date_cell.alignment = center
    date_cell.border    = b_all
    ws1.row_dimensions[1].height = 26

    # 下拉数据验证
    date_list_str = ",".join(str(d) for d in all_dates)
    dv = DataValidation(type="list", formula1=f'"{date_list_str}"', allow_blank=True)
    dv.sqref = "B1"
    ws1.add_data_validation(dv)

    sep_col = 10   # 均价区起始（1-based）

    def _sumproduct_formula(freeze: int, rank: int, hist_col: str) -> str:
        """生成 SUMPRODUCT+INDEX 公式，从 历史数据 sheet 中查询对应单元格值。"""
        cond = (
            f"(历史数据!$A$2:$A$5000=$B$1)"
            f"*(历史数据!$B$2:$B$5000={freeze})"
            f"*(历史数据!$C$2:$C$5000={rank})"
        )
        return (
            f'=IFERROR(IF(SUMPRODUCT({cond})=0,"",'
            f'INDEX(历史数据!${hist_col}:${hist_col},'
            f'SUMPRODUCT({cond}*ROW(历史数据!$A$2:$A$5000)))),"")'
        )

    def write_section_formulas(start_row: int, freeze: int, title: str) -> int:
        c_title  = CLR_30_TITLE  if freeze == 30 else CLR_7_TITLE
        c_header = CLR_30_HEADER if freeze == 30 else CLR_7_HEADER
        c_row_a  = CLR_30_ROW_A  if freeze == 30 else CLR_7_ROW_A
        c_row_b  = CLR_30_ROW_B  if freeze == 30 else CLR_7_ROW_B

        # 标题行
        ws1.merge_cells(start_row=start_row, start_column=1,
                        end_row=start_row, end_column=n_disp)
        tc = ws1.cell(row=start_row, column=1, value=title)
        tc.fill = fill(c_title)
        tc.font = Font(name="微软雅黑", bold=True, color=CLR_FONT_WHITE, size=12)
        tc.alignment = center
        tc.border = b_all
        ws1.row_dimensions[start_row].height = 22

        # 表头行
        hdr_row = start_row + 1
        for ci, h in enumerate(_DISP_HEADERS, start=1):
            c = ws1.cell(row=hdr_row, column=ci, value=h)
            c.fill = fill(c_header)
            c.font = Font(name="微软雅黑", bold=True,
                          color=CLR_FONT_WHITE if freeze == 7 else CLR_FONT_DARK,
                          size=10)
            c.alignment = center
            c.border = b_sep if ci == sep_col else b_all
        ws1.row_dimensions[hdr_row].height = 18

        # 数据行（公式）
        for rank in range(1, 6):
            dr = hdr_row + rank
            rf = fill(c_row_a if rank % 2 == 1 else c_row_b)
            ws1.row_dimensions[dr].height = 17
            for ci, (hist_col, disp_hdr) in enumerate(
                    zip(_HIST_DATA_COLS, _DISP_HEADERS), start=1):
                formula = _sumproduct_formula(freeze, rank, hist_col)
                c = ws1.cell(row=dr, column=ci, value=formula)
                c.fill      = rf
                c.border    = b_sep if ci == sep_col else b_all
                c.alignment = left_ if ci <= 2 else center
                if disp_hdr in ("税后利润", "税后利润(均)"):
                    c.number_format = '#,##0'
                    c.font = myfont()
                elif disp_hdr in ("收益率%", "收益率%(均)"):
                    c.number_format = '0.0"%"'
                    c.font = myfont()
                elif disp_hdr in ("买入价(最低)", "卖出价(最低)", "买入价(均价)", "卖出价(均价)"):
                    c.number_format = '#,##0'
                    c.font = myfont()
                else:
                    c.font = myfont()

        return hdr_row + 6  # 下一个可用行

    row_cursor = 3   # 第 2 行空出，从第 3 行开始写第一个分组
    row_cursor = write_section_formulas(
        row_cursor, 30, "▶ 高级兽诀 / 高级内丹（冻结 30 天）")
    row_cursor += 1  # 空一行
    write_section_formulas(
        row_cursor, 7,  "▶ 其余商品（冻结 7 天）")

    ws1.freeze_panes = "A2"

    # ════════════════════════════════════════════════════════════════════════
    # Sheet 2: 历史数据
    # ════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("历史数据")

    HDR_BG = "2E4057"
    HIST_COL_WIDTHS = [10, 8, 5, 12, 10, 8, 8, 7, 13, 13, 11, 9, 13, 13, 11, 9]
    for i, w in enumerate(HIST_COL_WIDTHS, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    for ci, h in enumerate(_HIST_HEADERS, start=1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.fill = fill(HDR_BG)
        c.font = Font(name="微软雅黑", bold=True, color=CLR_FONT_WHITE, size=10)
        c.alignment = center
        c.border = b_all
    ws2.row_dimensions[1].height = 20

    stripe_a = fill("F7F9FC")
    stripe_b = fill("FFFFFF")
    for ri, row in enumerate(rows, start=2):
        rf = stripe_a if ri % 2 == 0 else stripe_b
        ws2.row_dimensions[ri].height = 16
        for ci, h in enumerate(_HIST_HEADERS, start=1):
            val = row.get(h, "")
            c = ws2.cell(row=ri, column=ci, value=val)
            c.fill      = rf
            c.border    = b_all
            c.alignment = left_ if h == "商品名" else center
            if h in ("税后利润", "税后利润(均)"):
                if isinstance(val, (int, float)):
                    c.font = Font(name="微软雅黑", bold=True, size=10,
                                  color=CLR_PROFIT_POS if val >= 0 else CLR_PROFIT_NEG)
                else:
                    c.font = myfont()
                c.number_format = '#,##0'
            elif h in ("收益率%", "收益率%(均)"):
                if isinstance(val, (int, float)):
                    c.font = Font(name="微软雅黑", bold=True, size=10,
                                  color=CLR_PROFIT_POS if val >= 0 else CLR_PROFIT_NEG)
                else:
                    c.font = myfont()
                c.number_format = '0.0"%"'
            elif h in ("买入价(最低)", "卖出价(最低)", "买入价(均价)", "卖出价(均价)"):
                c.number_format = '#,##0'
                c.font = myfont()
            else:
                c.font = myfont()

    ws2.freeze_panes = "A2"

    os.makedirs(ANALYSIS_ROOT, exist_ok=True)
    wb.save(HISTORY_XLSX)
    print(f"历史汇总已保存：analysis/top_opportunities_history.xlsx"
          f"（{len(all_dates)} 天，共 {len(rows)} 行）")
